from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path("/mnt/c/shuangdian/05-1 作品信息概要表（人工智能实践赛、挑战赛，2026版）模板.docx")
OUTPUT = ROOT / "docs" / "competition" / "2026002224-作品信息概要表.docx"
README_OUTPUT = ROOT / "docs" / "competition" / "2026002224-03设计与开发文档-readme.txt"


def set_run_font(run, size: float = 9.5, bold: bool = False) -> None:
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.font.bold = bold


def set_cell_text(cell, text: str, *, size: float = 9.5, bold: bool = False, align=None) -> None:
    cell.text = ""
    parts = text.split("\n")
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.LEFT
    for index, part in enumerate(parts):
        if index:
            paragraph = cell.add_paragraph()
            paragraph.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.LEFT
        run = paragraph.add_run(part)
        set_run_font(run, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_document(document: Document) -> None:
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            set_run_font(run, 10.5)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        set_run_font(run, 9.5)


def fill_basic_info(table) -> None:
    set_cell_text(table.cell(0, 3), "2026002224", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(table.cell(0, 10), "牧野智农——智能农业害虫防治系统", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(table.cell(1, 12), "■实践赛  □挑战赛", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    intro = (
        "作品简介(100字以内)：\n"
        "牧野智农面向大田植保场景，集成 YOLO 害虫识别、天气融合、多模型AI会诊（Qwen/DeepSeek/Xiaomi）、合规推理链、药效评估闭环和 React 指挥大屏，形成从虫情感知到无人机执行再到药效验证的完整自动化闭环系统。"
    )
    innovation = (
        "创新描述（100字以内）：\n"
        "（1）多模型加权投票降低单一AI偏差；（2）五维合规推理链保障施药安全；（3）药效评估闭环实现自动补充喷洒；（4）变量喷洒按虫害密度精准施药；（5）仿真与真实无人机双模执行，灵活适配教学与生产场景。"
    )
    special = (
        "特别说明（100字以内）：\n"
        "地图为前端 SVG 虚拟农田示意，不接入真实地图底图。系统支持真实无人机与仿真双模执行，可一键启停。使用 AI 工具辅助代码分析、文档草稿和调试建议，内容经团队核验，核心设计与实现由团队完成。"
    )
    set_cell_text(table.cell(2, 0), intro, size=9)
    set_cell_text(table.cell(3, 0), innovation, size=9)
    set_cell_text(table.cell(4, 0), special, size=9)


def fill_authors(table) -> None:
    # 参赛成员姓名和分工比例需要以最终报名信息为准，这里只保留待补充占位。
    author_headers = ["待补充1", "待补充2", "待补充3", "待补充4", "待补充5"]
    for col, name in zip([2, 6, 9, 11, 14], author_headers):
        set_cell_text(table.cell(6, col), name, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    for row in range(7, 14):
        for col in [2, 6, 9, 11, 14]:
            set_cell_text(table.cell(row, col), "待补充", align=WD_ALIGN_PARAGRAPH.CENTER)

    teacher_roles = (
        "■项目创意  ■理论指导  ■技术方案  □实验场地  □硬件资源\n"
        "□数据提供  □后勤支持  ■宣讲通知  ■组织协调  □经费支持\n"
        "□其他：答辩材料修改建议"
    )
    set_cell_text(table.cell(14, 5), teacher_roles, size=9)


def fill_platforms(table) -> None:
    set_cell_text(table.cell(15, 4), "■Windows  ■Linux  □macOS  ■其他：WSL/Ubuntu、PX4/Gazebo 仿真环境", size=9)
    set_cell_text(table.cell(16, 4), "■Windows  ■Linux  □macOS  □iOS  □Android  ■其他：Chrome/Edge 浏览器、PX4 SITL/Gazebo", size=9)
    tools = (
        "Python 3.11、FastAPI、React/Vite/TypeScript、YOLO、SQLite、ChromaDB、MAVSDK、PX4 SITL、Gazebo；\n"
        "后台大语言模型：通义千问（Qwen，DashScope）、DeepSeek（DeepSeek Chat）、小米大模型（Xiaomi MiMo）；\n"
        "文本向量模型：text-embedding-v3。"
    )
    set_cell_text(table.cell(17, 4), tools, size=8.8)
    refs = (
        "1、PX4 与 MAVSDK 官方文档\n"
        "2、FastAPI、React、Vite 官方文档\n"
        "3、Ultralytics YOLO、ChromaDB、DashScope、和风天气开发文档"
    )
    set_cell_text(table.cell(18, 4), refs, size=8.8)
    submits = "■报告文档  ■展示视频  ■PPT  ■源代码  ■部署文件  ■数据集  ■模型\n■其他：readme.txt、运行说明、配置说明"
    set_cell_text(table.cell(19, 4), submits, size=9)


def fill_related_files(table) -> None:
    rows = [
        (
            "文件：2026002224-作品信息概要表.docx\n描述：作品基础信息、简介、创新点、开发平台、提交内容及相关文件说明。",
            "■已上传到网盘\n□未上传，下载地址：随参赛总文件夹提交",
            "■自制  □未知版权\n□开源  □授权方：",
        ),
        (
            "文件：2026002224-作品报告.docx\n描述：项目背景、需求分析、技术方案（含多模型AI会诊、合规推理链、药效评估闭环）、系统实现、测试分析和作品总结。",
            "■已上传到网盘\n□未上传，下载地址：随参赛总文件夹提交",
            "■自制  □未知版权\n□开源  □授权方：",
        ),
        (
            "文件：muye 项目源码\n描述：FastAPI 后端、YOLO 接口、多模型会诊决策（Qwen/DeepSeek/Xiaomi）、合规推理链、药效评估、PX4/真实无人机双模执行、React 前端、测试与部署脚本。",
            "■已上传到网盘\n□未上传，下载地址：随参赛总文件夹提交",
            "■自制  □未知版权\n■开源组件  授权方：见依赖许可证",
        ),
        (
            "文件：data/samples 与 data/seeds\n描述：展示图片、农业种子数据、农药目录、地块和作物周期等测试数据。",
            "■已上传到网盘\n□未上传，下载地址：随参赛总文件夹提交",
            "■自制整理  □未知版权\n□开源  □授权方：",
        ),
        (
            "文件：YOLO 检测模型与配置\n描述：害虫检测模型权重、类别配置和本地推理服务配置。",
            "■已上传到网盘\n□未上传，下载地址：随参赛总文件夹提交",
            "■自制/训练  □未知版权\n■开源框架  授权方：Ultralytics YOLO",
        ),
        (
            "文件：readme.txt\n描述：说明本文件夹用途，并逐项说明设计与开发文档目录中的文件。",
            "■已上传到网盘\n□未上传，下载地址：随参赛总文件夹提交",
            "■自制  □未知版权\n□开源  □授权方：",
        ),
    ]
    for row_index, (desc, status, copyright_status) in enumerate(rows, start=22):
        set_cell_text(table.cell(row_index, 1), desc, size=8.5)
        set_cell_text(table.cell(row_index, 8), status, size=8.5)
        set_cell_text(table.cell(row_index, 13), copyright_status, size=8.5)


def write_readme() -> None:
    content = (
        "2026002224-03\u8bbe\u8ba1\u4e0e\u5f00\u53d1\u6587\u6863\n"
        "\n"
        "\u672c\u6587\u4ef6\u5939\u4f5c\u7528\uff1a\n"
        "\u7528\u4e8e\u5b58\u653e\u7267\u91ce\u667a\u519c\u2014\u2014\u667a\u80fd\u519c\u4e1a\u5bb3\u866b\u9632\u6cbb\u7cfb\u7edf\u7684\u8bbe\u8ba1\u3001\u5f00\u53d1\u3001\u8bf4\u660e\u7c7b\u6587\u6863\uff0c\u4f9b\u5927\u8d5b\u8bc4\u5ba1\u4e86\u89e3\u4f5c\u54c1\u7684\u603b\u4f53\u65b9\u6848\u3001\u6280\u672f\u8def\u7ebf\uff08\u542b\u591a\u6a21\u578bAI\u4f1a\u8bca\u3001\u5408\u89c4\u63a8\u7406\u94fe\u3001\u836f\u6548\u8bc4\u4f30\u95ed\u73af\uff09\u3001\u5b9e\u73b0\u8fc7\u7a0b\u3001\u6d4b\u8bd5\u60c5\u51b5\u548c\u76f8\u5173\u6587\u4ef6\u4fe1\u606f\u3002\n"
        "\n"
        "\u6587\u4ef6\u8bf4\u660e\uff1a\n"
        "1. 2026002224-\u4f5c\u54c1\u62a5\u544a.docx\n"
        "   \u57fa\u4e8e\u4eba\u5de5\u667a\u80fd\u5b9e\u8df5\u8d5b\u4f5c\u54c1\u62a5\u544a\u6a21\u677f\u7f16\u5199\uff0c\u5185\u5bb9\u5305\u62ec\u4f5c\u54c1\u6982\u8ff0\u3001\u95ee\u9898\u5206\u6790\u3001\u89e3\u51b3\u601d\u8def\u3001\u6280\u672f\u65b9\u6848\uff08\u591a\u667a\u80fd\u4f53\u4f1a\u8bca\u3001\u51b3\u7b56\u8def\u7531\u3001\u5408\u89c4\u63a8\u7406\u94fe\u3001\u53d8\u91cf\u55b7\u6d12\u3001\u836f\u6548\u8bc4\u4f30\u95ed\u73af\uff09\u3001\u7cfb\u7edf\u5b9e\u73b0\u3001\u6d4b\u8bd5\u5206\u6790\u3001\u4f5c\u54c1\u603b\u7ed3\u3001\u521b\u65b0\u70b9\u3001\u5e94\u7528\u63a8\u5e7f\u548c\u53c2\u8003\u8d44\u6599\u3002\n"
        "\n"
        "2. 2026002224-\u4f5c\u54c1\u4fe1\u606f\u6982\u8981\u8868.docx\n"
        "   \u57fa\u4e8e\u4eba\u5de5\u667a\u80fd\u5b9e\u8df5\u8d5b/\u6311\u6218\u8d5b\u4f5c\u54c1\u4fe1\u606f\u6982\u8981\u8868\u6a21\u677f\u586b\u5199\uff0c\u5185\u5bb9\u5305\u62ec\u4f5c\u54c1\u7f16\u53f7\u3001\u4f5c\u54c1\u540d\u79f0\u3001\u4f5c\u54c1\u7b80\u4ecb\u3001\u521b\u65b0\u63cf\u8ff0\u3001\u5f00\u53d1\u5236\u4f5c\u5e73\u53f0\u3001\u8fd0\u884c\u5c55\u793a\u5e73\u53f0\u3001\u5f00\u53d1\u5de5\u5177\u3001\u63d0\u4ea4\u5185\u5bb9\u548c\u76f8\u5173\u6587\u4ef6\u8bf4\u660e\u3002\n"
        "\n"
        "3. \u4e2d\u56fd\u5927\u5b66\u751f\u8ba1\u7b97\u673a\u8bbe\u8ba1\u5927\u8d5bAI\u5de5\u5177\u4f7f\u7528\u8bf4\u660e.pdf\n"
        "   \u5927\u8d5b\u76f8\u5173\u8bf4\u660e\u6587\u6863\uff0c\u7528\u4e8e\u6838\u5bf9\u4eba\u5de5\u667a\u80fd\u5de5\u5177\u4f7f\u7528\u3001\u53c2\u8d5b\u6750\u6599\u586b\u5199\u548c\u63d0\u4ea4\u89c4\u8303\u3002\n"
        "\n"
        "4. readme.txt\n"
        "   \u672c\u8bf4\u660e\u6587\u4ef6\uff0c\u7b80\u8981\u8bf4\u660e\u5f53\u524d\u6587\u4ef6\u5939\u4f5c\u7528\uff0c\u5e76\u5bf9\u6587\u4ef6\u5939\u5185\u5404\u6587\u4ef6\u8fdb\u884c\u63cf\u8ff0\u3002\n"
    )
    README_OUTPUT.write_text(content, encoding="utf-8")


def main() -> None:
    if not TEMPLATE.exists():
        raise FileNotFoundError(TEMPLATE)
    document = Document(TEMPLATE)
    style_document(document)
    table = document.tables[0]
    fill_basic_info(table)
    fill_authors(table)
    fill_platforms(table)
    fill_related_files(table)
    document.save(OUTPUT)
    write_readme()
    print(OUTPUT)
    print(README_OUTPUT)


if __name__ == "__main__":
    main()
