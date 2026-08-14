from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path("/mnt/c/shuangdian/05-3 作品报告（人工智能实践赛，2026版）模板.docx")
OUTPUT = ROOT / "docs" / "competition" / "2026002224-作品报告.docx"


def set_run_font(run, size: int | float | None = None, bold: bool | None = None) -> None:
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold


def set_para_font(paragraph, size: int | float | None = None, bold: bool | None = None) -> None:
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold)


def clear_document(document: Document) -> None:
    body = document._body._element
    for child in list(body):
        if child.tag.endswith("}sectPr"):
            continue
        body.remove(child)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.7)
    section.right_margin = Cm(2.7)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
        style = styles[style_name]
        style.font.name = "宋体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)


def paragraph(document: Document, text: str = "", *, first_line: bool = True, align=None):
    p = document.add_paragraph()
    if align is not None:
        p.alignment = align
    if first_line:
        p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, 11)
    return p


def heading(document: Document, text: str, level: int = 1):
    p = document.add_heading(text, level=level)
    set_para_font(p, {1: 16, 2: 14, 3: 12}.get(level, 11), True)
    p.paragraph_format.first_line_indent = Cm(0)
    return p


def bullet(document: Document, items: list[str]) -> None:
    for index, item in enumerate(items, start=1):
        p = document.add_paragraph()
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        run = p.add_run(f"（{index}）{item}")
        set_run_font(run, 11)


def table(document: Document, caption: str, headers: list[str], rows: list[list[str]]) -> None:
    p = paragraph(document, caption, first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_para_font(p, 10.5, True)
    tbl = document.add_table(rows=1, cols=len(headers))
    tbl.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = tbl.rows[0].cells[idx]
        cell.text = header
        for run in cell.paragraphs[0].runs:
            set_run_font(run, 10.5, True)
    for row in rows:
        cells = tbl.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
            for para in cells[idx].paragraphs:
                for run in para.runs:
                    set_run_font(run, 10)
    document.add_paragraph()


def flow_box(document: Document) -> None:
    caption = paragraph(document, "图 1  系统技术路线框架图", first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_para_font(caption, 10.5, True)
    tbl = document.add_table(rows=9, cols=1)
    tbl.style = "Table Grid"
    steps = [
        "巡检图像采集 / 图片上传",
        "YOLO 害虫识别：输出害虫类别、置信度与标注图",
        "天气与地块上下文融合：温度、湿度、风速、作物与地块参数",
        "RAG 知识增强：农药目录、历史案例、农业知识片段检索",
        "多智能体AI会诊：Qwen/DeepSeek/Xiaomi三模型加权投票 + 合规推理链",
        "结构化决策输出：生成用药方案、农事建议、安全提示与合规结论",
        "任务规划与无人机执行：变速率喷洒航线、人工确认、Mission执行",
        "药效评估闭环：复检图片YOLO识别 → 杀灭率计算 → 自动决策补充喷洒",
        "React 指挥大屏：地图动画、流程状态、决策可视化、失败原因、历史追溯",
    ]
    for row, step in zip(tbl.rows, steps):
        cell = row.cells[0]
        cell.text = step
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in para.runs:
            set_run_font(run, 10.5)
            if "：" in step:
                run.font.color.rgb = RGBColor(31, 78, 121)
    document.add_paragraph()


def add_cover(document: Document) -> None:
    for text, size, bold in [
        ("2026年（第19届）", 18, True),
        ("中国大学生计算机设计大赛", 22, True),
        ("人工智能实践赛作品报告", 20, True),
    ]:
        p = paragraph(document, text, first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER)
        p.paragraph_format.space_after = Pt(12)
        set_para_font(p, size, bold)

    document.add_paragraph()
    for text in [
        "作品编号：2026002224",
        "作品名称：牧野智农——智能农业害虫防治系统",
        "填写日期：2026年5月14日",
    ]:
        p = paragraph(document, text, first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_para_font(p, 14, True)
        p.paragraph_format.space_after = Pt(10)
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_report(document: Document) -> None:
    add_cover(document)

    heading(document, "一、作品概述", 1)
    paragraph(
        document,
        "《牧野智农——智能农业害虫防治系统》面向大田植保作业场景，围绕“虫情感知、气象感知、AI 决策、无人机执行、前端可视化”构建端到端闭环。作品不是单独展示某个模型，而是把图像识别、天气数据、知识增强决策、无人机仿真执行和指挥大屏串联成可运行、可追溯、可扩展的智慧农业原型系统。",
    )
    paragraph(
        document,
        "系统支持用户上传巡检图片或由采集模块投喂样本图片；后端调用本地 YOLO 服务识别害虫类型、置信度和位置；天气模块融合实时或模拟气象数据；决策模块通过多智能体会诊（Qwen、DeepSeek、Xiaomi 三模型加权投票）、RAG 知识增强和合规推理链生成结构化施药建议；无人机模块把决策结果转化为变速率喷洒航线，支持仿真与真实无人机执行；药效评估模块通过复检图片计算杀灭率并自动决定是否需要补充喷洒；前端大屏展示地图动画、任务流程、会诊详情、合规结论、失败原因和历史记录。",
    )
    bullet(
        document,
        [
            "服务对象：植保服务队、农场与合作社、农业物联网平台、智慧农业教学与科研验证场景。",
            "核心价值：将“人工巡田、经验判断、手动飞行”的离散流程升级为“数据采集、智能决策、仿真执行、过程留痕”的闭环系统。",
            "展示特点：系统可在本地一键启动，支持 FastAPI 主服务、YOLO 推理服务、React 指挥大屏和 PX4 SITL 链路联调。",
        ],
    )

    heading(document, "二、问题分析", 1)
    heading(document, "2.1 问题来源", 2)
    paragraph(
        document,
        "农作物病虫害具有传播快、局部爆发强、治理窗口短的特点。传统大田植保依赖人工巡田、经验识别和飞手手动规划航线，容易出现发现不及时、施药不精准、用药记录缺失、作业效果难以追溯等问题。随着种植规模扩大，仅依靠人工经验已经难以同时满足效率、精度和安全要求。",
    )
    paragraph(
        document,
        "现实作业中，图像识别工具、天气数据、农药知识库、无人机飞控软件和生产记录通常分散在不同系统中。即使单个环节已有成熟工具，仍然缺少一套能够把“检测结果如何变成施药方案、施药方案如何变成飞行任务、执行状态如何回到管理端”的完整链路。牧野项目正是针对这一断层进行工程化验证。",
    )

    heading(document, "2.2 现有解决方案", 2)
    table(
        document,
        "表 1  现有方案与本作品对比",
        ["方案类型", "主要能力", "不足", "牧野改进点"],
        [
            ["人工巡田与经验施药", "人工观察虫情，依据经验配药和安排作业", "覆盖慢、判断主观、缺少数据记录", "用 YOLO 识别和结构化流程替代纯经验输入"],
            ["单点害虫识别应用", "识别图片中的害虫类别", "通常止步于识别结果，不能形成执行任务", "把识别结果继续送入天气、RAG、决策和无人机规划"],
            ["通用植保无人机软件", "支持航线规划和飞行控制", "对虫情识别、用药决策、知识检索支持不足", "将 AI 决策与 PX4 SITL 执行链路打通"],
            ["规则型农事专家系统", "按规则给出农事建议", "适配复杂场景能力有限，维护成本高", "引入 Qwen 大模型和 ChromaDB 知识检索并保留结构化校验"],
        ],
    )

    heading(document, "2.3 本作品要解决的痛点问题", 2)
    bullet(
        document,
        [
            "识别与决策割裂：害虫检测结果不能直接支撑施药方案和无人机作业参数。",
            "决策上下文不足：施药建议往往缺少天气、作物、农药知识和历史案例支撑，单一模型判断可能存在偏差。",
            "执行链路缺失：AI 输出无法稳定转化为可验证的航线、高度、速度和喷洒任务。",
            "合规安全被忽视：缺乏结构化的农药合规检查（来源验证、作物匹配、毒性评估、天气约束等），施药方案存在安全隐患。",
            "缺乏药效验证闭环：施药后没有自动化的效果评估与补充喷洒机制，虫害可能未得到充分灭杀。",
            "失败原因不可见：业务链路出错时，如果前端只显示失败状态，操作者无法快速判断是天气、识别、决策还是无人机链路异常。",
        ],
    )

    heading(document, "三、解决问题的思路", 1)
    paragraph(
        document,
        "作品采用“闭环优先、模块解耦、可降级运行”的总体思路。系统先保证从图片输入到前端展示的主流程可以稳定跑通，再逐步接入外部天气、大模型、RAG 向量库和 PX4 SITL。各模块之间通过 Pydantic 数据模型、事件总线和 SQLite 结构化存储传递状态，避免一个模块失败导致全系统不可用。",
    )
    heading(document, "3.1 需求分析", 2)
    bullet(
        document,
        [
            "功能需求：图片上传与采集、害虫识别、天气获取、知识增强决策、无人机任务规划、人工确认起飞、地图态势展示、任务历史追溯。",
            "性能需求：前端大屏应能持续刷新，后端任务状态应及时反馈，YOLO 和外部 API 调用失败时要有明确降级或错误提示。",
            "安全需求：无人机起飞默认保留人工确认，PX4 仿真使用本地坐标系执行航线，不向飞控上传业务地块的真实经纬度航点。",
            "可维护需求：项目按照 app、modules、models、data、config、frontend 分层，便于后续替换模型、扩展接口和接入真实设备。",
        ],
    )
    heading(document, "3.2 数据来源与样本", 2)
    paragraph(
        document,
        "项目数据由展示样本、配置数据、运行数据和外部服务数据共同构成。展示样本位于 data/samples，覆盖蚜虫、稻飞虱、玉米螟、蛴螬、红蜘蛛等害虫图片；农业种子数据位于 data/seeds/henan，包含地块、作物周期、土壤记录、农药目录和统计指标；运行数据由 JSONL 事件流和 SQLite 数据库记录；天气数据可通过和风天气 API 获取，也可以在无密钥场景下降级为 mock 数据。",
    )
    bullet(
        document,
        [
            "检测数据字段：pest_type、confidence、position、original_image、annotated_image。",
            "天气数据字段：temperature、humidity、wind_speed、wind_direction、precipitation、condition。",
            "决策数据字段：药剂建议、浓度、用量、施药窗口、风险提示、农事建议、RAG 上下文摘要。",
            "无人机数据字段：飞行路径、高度、速度、喷洒速率、覆盖区域、任务阶段、电量与位置状态。",
        ],
    )
    heading(document, "3.3 总体实现思路", 2)
    paragraph(
        document,
        "系统将“检测、会诊、决策、合规、执行、评估、展示”拆分为独立领域模块。检测域只负责图片处理和害虫结构化输出；会诊域汇聚多模型专家投票并路由决策路径；决策域根据害虫、天气和知识上下文生成施药建议；合规域通过多角度推理链验证方案安全性；任务规划器把建议转换为飞行路径和喷洒参数；无人机模块负责仿真或真实执行；评估模块通过复检计算杀灭率实现闭环；前端负责把复杂流程转化为可理解的可视化界面。这样的职责边界使系统既能展示完整闭环，又能在单个模块替换时保持整体稳定。",
    )

    heading(document, "四、技术方案", 1)
    flow_box(document)
    heading(document, "4.1 总体架构", 2)
    paragraph(
        document,
        "系统采用前后端分离、领域模块化和混合存储架构。前端使用 React、Vite、TypeScript 构建指挥大屏；应用层使用 FastAPI 提供 REST API 与 WebSocket；核心业务层按 detection、decision、drone、infra 四个领域组织；模型层使用 Pydantic 统一接口结构；数据层使用 SQLite、JSONL 事件流、ChromaDB 和配置文件共同支撑运行。",
    )
    table(
        document,
        "表 2  系统主要模块",
        ["模块", "关键技术", "职责"],
        [
            ["前端展示层", "React + Vite + TypeScript + SVG", "展示地图动画、流程状态、会诊详情、合规结论、失败原因和历史记录"],
            ["应用编排层", "FastAPI + asyncio + Pydantic", "组织 API、任务队列、WebSocket 推送和主流程编排"],
            ["害虫检测域", "YOLO 本地推理服务", "识别害虫类别、置信度、位置并生成标注结果"],
            ["智能决策域", "Qwen + DeepSeek + Xiaomi + RAG + ChromaDB + jsonschema", "多模型加权投票、知识增强检索、熟悉度路由、合规推理链"],
            ["无人机执行域", "PX4 SITL / 真实无人机 + MAVSDK Mission", "支持仿真和真实飞控两种执行模式，变量喷洒路径规划"],
            ["药效评估域", "YOLO 复检 + SQLite 评估存储", "施药后复检图片识别、杀灭率计算、自动触发补充喷洒"],
            ["基础设施域", "SQLite + JSONL EventBus + Weather API", "记录任务、事件、天气、决策、评估和无人机状态"],
        ],
    )
    heading(document, "4.2 关键算法与协议", 2)
    bullet(
        document,
        [
            "害虫识别：通过 YOLO 模型服务输出检测框和类别置信度，主流程只消费统一的结构化结果，便于模型替换。",
            "知识增强：使用 text-embedding-v3 生成向量，ChromaDB 管理 pesticides、decisions、agri_knowledge 三类 collection，检索结果拼入决策上下文。",
            "多智能体加权投票：系统内置昆虫学家、农学家、农药专家三个角色，通过 Qwen、DeepSeek、Xiaomi 三个不同模型生成独立用药建议，基于熟悉度评分计算权重并加权投票选出最优方案。",
            "决策路由（DecisionRouter）：根据农药匹配度、历史案例相似度、作物匹配度和目录覆盖率四个维度计算熟悉度分数，自动在单专家快速路由和多智能体会诊之间切换，平衡效率与决策质量。",
            "合规推理链（PesticideComplianceChecker）：对施药方案进行来源验证、作物匹配、毒性评估、天气约束和安全提示五个维度检查，每项输出 passed/warning/blocked 状态及证据溯源，生成执行策略（自动/手动/阻止）并推荐替代方案。",
            "结构化决策：多模型会诊结果经 jsonschema 校验后进入后续链路，避免自然语言结果直接驱动无人机任务。",
            "变量喷洒航线规划：任务规划器根据 YOLO 检测密度分布生成密度热力图，计算逐格喷洒计划（spray_schedule），实现按虫害密度变量施药，减少农药浪费。",
            "无人机执行：系统采用 MAVSDK Mission 插件，将展示米制航线锚定到当前 Home 附近后交由 Mission 模式执行；支持仿真（PX4 SITL）与真实无人机两种执行模式，不使用业务地块经纬度作为飞控航点。",
            "药效评估闭环：施药后上传复检图片，再次调用 YOLO 识别虫口数，计算杀灭率 =（施药前 - 施药后）/ 施药前 × 100%，低于阈值（默认 90%）时自动触发补充喷洒，最多迭代 3 次。",
            "状态同步：后端通过 HTTP 与 WebSocket 推送 workflow_state 和 sim_map，前端断线后降级为轮询，保证大屏不断档。",
        ],
    )
    heading(document, "4.3 技术特色", 2)
    paragraph(
        document,
        "作品的技术特色在于把多模型AI会诊、合规推理、变量喷洒和药效评估等环节全部放入同一条可运行链路，而不是把识别、问答和无人机展示做成孤立页面。系统明确区分“AI 会诊决策”、“合规验证”和“任务规划”职责，既发挥大模型的农事建议生成能力，又用合规推理链和规则化规划器保证施药方案安全可控。多模型加权投票机制降低了单一模型判断偏差，熟悉度路由根据场景自动选择决策路径；药效评估闭环实现了从“识别→施药→评估→再施药”的完整自动化。同时，无人机链路支持仿真和真实飞控双模切换，解决演示与生产的衔接问题。",
    )

    heading(document, "五、系统实现", 1)
    heading(document, "5.1 工程结构", 2)
    paragraph(
        document,
        "项目根目录按照职责划分为 app、modules、models、frontend、data、config、scripts、tests 等部分。app 负责 FastAPI 路由和服务编排；modules 承载核心业务；models 定义跨模块数据结构；frontend 是当前唯一前端主线；data 存放样本、运行事件、SQLite 数据和种子数据；config 保存无人机、YOLO 等配置；scripts 提供准备、检查和一键展示脚本。",
    )
    bullet(
        document,
        [
            "app/routes：health、workflow、dashboard、demo、evaluation 等接口按领域拆分。",
            "app/services：聚合工作流状态、地图状态、评估管理、演示数据播种等逻辑。",
            "modules/detection：图像处理、本地 YOLO API、数据采集。",
            "modules/decision：AI 决策核心——包含 agents（多智能体会诊、加权投票）、rag（向量检索、知识加载）、router（熟悉度路由）、compliance（合规推理链）。",
            "modules/drone：任务规划、PX4 仿真控制、地块上下文解析、变量喷洒路径。",
            "modules/infra：事件总线、SQLite 存储、天气服务和公共工具。",
        ],
    )
    heading(document, "5.2 主流程实现", 2)
    paragraph(
        document,
        "用户在前端上传图片后，主 API 将图片加入任务队列。后台 worker 调用 YOLO 服务完成检测，并把检测结果、天气信息和地块上下文交给决策模块。决策阶段首先通过 DecisionRouter 计算熟悉度分数，决定走单专家快速路由或多智能体会诊；ExpertConsultation 调用昆虫学家、农学家、农药专家三个角色通过不同 LLM 生成独立建议，经加权投票选出最优方案；PesticideComplianceChecker 对方案进行五维合规检查并输出执行策略。决策完成后，系统生成变速率喷洒航线并进入人工确认状态。用户确认起飞后，无人机执行 Mission，施药后自动触发复检评估，杀灭率不达标时自动补充喷洒。所有阶段通过事件总线和 SQLite 记录，前端实时展示流程进度。",
    )
    paragraph(
        document,
        "为了保障演示的灵活性和展示稳定性，系统支持演示模式与真实作业模式双轨运行。演示模式通过 seed_demo_stages.py 预播种数据，前端展示完整的五阶段作业过程（巡检→检测→会诊→喷洒→评估），无需连接真实无人机和外部 API；真实模式则接入实际 YOLO 推理、天气服务和飞控链路。前端地图大屏以任务规划航线和前端动画展示作业态势，避免飞控状态抖动影响大屏观感。",
    )
    heading(document, "5.3 关键功能实现", 2)
    bullet(
        document,
        [
            "一键演示：scripts/demo.sh 可启动后端 API 和 React 前端，scripts/seed_demo_stages.py 支持分阶段注入演示数据，docs/competition/demo-script.md 提供配套解说词，三件套支撑完整竞赛演示。",
            "人工确认起飞：后端在 pending_confirmation 阶段等待用户确认，前端提供确认按钮，避免无人机任务自动越过安全确认。",
            "多模型加权投票：ExpertConsultation 内置三个专家角色，支持配置不同 LLM provider，通过熟悉度加权投票产生最终用药方案。",
            "合规推理链：PesticideComplianceChecker 对施药方案执行五维检查（来源验证、作物匹配、毒性评估、天气约束、安全提示），每项附带证据溯源和应对建议，输出结构化执行策略。",
            "药效评估闭环：施药后上传复检图片，YOLO 重新识别虫口数，计算杀灭率，低于阈值时自动创建新迭代补充喷洒，支持配置最大迭代次数。",
            "变量喷洒：任务规划器根据 YOLO 检测位置生成密度热力图（density_grid），逐格计算喷洒需求（spray_schedule），实现按需施药。",
            "业务失败原因展示：workflow 事件 payload 向前端透传错误细节，前端在流程面板中显示业务失败原因，便于定位天气、检测、决策或无人机异常。",
            "混合存储：JSONL 负责实时事件流和时间线回放，SQLite 负责结构化任务摘要、检测、天气、决策、合规、评估、无人机状态和农业基础数据。",
            "RAG 降级：知识检索失败时记录告警并回退基础决策，不阻断检测和展示主链路。",
        ],
    )
    heading(document, "5.4 运行与部署", 2)
    paragraph(
        document,
        "系统默认运行端口为：主 API 服务 18000，YOLO 推理服务 8010，前端大屏 5173。prepare.sh 用于检查依赖、样本图片和 RAG 知识库；demo.sh 是一键展示入口；precheck.sh 提供可复用环境检查。项目也提供 Dockerfile、docker-compose 和 systemd/nginx 部署文件，便于后续工程化落地。",
    )

    heading(document, "六、测试分析", 1)
    paragraph(
        document,
        "项目采用自动化测试、前端构建测试、接口联调和手工展示验证相结合的方式。核心链路覆盖应用层、决策域（AI决策、会诊、合规推理）、无人机域、检测域和基础设施域。近期联调中还针对演示流程、多模型会诊、合规推理链和药效评估闭环进行了专项验证。",
    )
    table(
        document,
        "表 3  测试与验证情况",
        ["测试类型", "覆盖对象", "验证重点", "结果"],
        [
            ["单元测试", "decision、agents、compliance、infra、drone 等模块", "配置解析、数据校验、RAG 检索、会诊投票、合规推理、SQLite CRUD、任务规划", "核心模块通过"],
            ["集成测试", "FastAPI 路由、workflow_service、SQLite 与事件总线、evaluation", "接口响应、任务历史、事件合并、状态聚合、评估记录", "主链路可运行"],
            ["前端构建", "frontend React 工程", "TypeScript 编译、Vite 生产构建、页面依赖完整性", "npm run build 通过"],
            ["无人机专项验证", "PX4Simulator、DroneController、MissionPlanner", "不上传经纬度航点，使用本地 NED 航线；双模执行", "相关验证通过"],
            ["演示验证", "scripts/demo.sh 全流程", "一键启动、五阶段数据播种、决策可视化、前端状态展示", "满足竞赛展示要求"],
        ],
    )
    paragraph(
        document,
        "测试发现并修正的关键问题包括：PX4 使用真实经纬度航点时与 SITL 本地地图不一致导致无人机未按预设航线飞行；单一 AI 模型判断存在农药推荐偏差；缺少合规检查导致施药方案安全评估不足；施药后缺少效果验证环节。针对这些问题，系统已分别改为本地航线执行、多模型加权投票、五维合规推理链和药效评估闭环机制。此外前端地图展示与右侧信息栏分离、工作流面板展示失败 payload 等改进也显著提升了展示可读性和调试效率。",
    )
    paragraph(
        document,
        "仍需继续加强的测试包括：多模型会诊的一致性与偏差测试、合规推理链的边界 case 测试、真实环境集成测试，以及更多不同虫害样本和异常天气组合下的端到端回归测试。",
    )

    heading(document, "七、作品总结", 1)
    heading(document, "7.1 作品特色与创新点", 2)
    bullet(
        document,
        [
            "全链路闭环：将害虫识别、天气融合、多智能体会诊、合规验证、无人机执行、药效评估和指挥大屏整合到同一系统。",
            "多模型AI协诊：内置昆虫学家、农学家、农药专家三个角色，通过 Qwen、DeepSeek、Xiaomi 三个不同模型加权投票，降低单一模型偏差。",
            "决策路由机制：基于熟悉度四维评分（农药匹配、历史案例、作物相似度、目录覆盖）自动在快速决策和多智能体会诊之间路由，兼顾效率与质量。",
            "合规推理链：五维施药安全审查（来源验证、作物匹配、毒性评估、天气约束、安全提示），每项附带证据溯源，输出结构化执行策略和替代方案。",
            "药效评估闭环：施药后自动复检、杀灭率计算、不达标则补充喷洒的完整反馈机制，实现真正的自动化防治。",
            "变频变量喷洒：根据虫害密度热力图计算逐格喷洒计划，实现按需施药，减少农药浪费。",
            "双模无人机执行：同一套航线规划支持仿真（PX4 SITL）和真实无人机两种执行模式，便于从演示过渡到生产。",
            "可视化与可追溯：React 大屏展示任务阶段、会诊推理、合规结论、地图动画、失败原因和历史记录，SQLite 与事件总线保留过程数据。",
            "工程化程度较高：交互式演示剧本、阶段化数据注入、配套解说词文档、模块化代码、自动化测试和部署文件为后续落地提供基础。",
        ],
    )
    heading(document, "7.2 应用推广", 2)
    paragraph(
        document,
        "作品可作为农业高校和科研团队的智慧植保教学平台，也可用于植保服务企业的无人机调度原型验证。多模型会诊和合规推理链已在演示模式中得到验证，可提供完整的智慧植保教学案例。对于农场和合作社，系统未来可接入真实巡检无人机、固定摄像头或田间传感器，把当前的演示模式扩展为生产环境中的虫情监测、施药决策辅助和药效验证平台。",
    )
    paragraph(
        document,
        "推广时可以分阶段落地：第一阶段部署本地识别和决策展示能力，服务于病虫害巡检和农事建议；第二阶段接入真实地块、作物周期、天气站和农药库存，激活合规推理和变量喷洒；第三阶段在安全许可下接入真实植保无人机和复检机制，实现从辅助决策到半自动执行再到药效验证的完整闭环。",
    )
    heading(document, "7.3 作品展望", 2)
    bullet(
        document,
        [
            "扩充害虫与病害数据集，提高模型在不同光照、作物和地区条件下的识别鲁棒性。",
            "引入多源传感器数据，如土壤墒情、叶面湿度、田间摄像头和历史虫情统计。",
            "完善药剂安全规则，增加农药安全间隔期、混配禁忌和环保约束，扩大合规检查维度。",
            "提升无人机从仿真到实飞的安全校验流程，增加空域申报、禁飞区规避等安全层。",
            "建设更完整的数据闭环，根据防治效果回写历史案例和熟悉度评分，持续优化 RAG 知识库和决策路由策略。",
            "扩展多智能体会诊角色和模型提供商，探索基于强化学习的决策权重自适应调整。",
        ],
    )

    heading(document, "八、参考文献", 1)
    references = [
        "Ultralytics YOLO Documentation.",
        "FastAPI Documentation.",
        "React Documentation.",
        "Vite Documentation.",
        "PX4 Autopilot User Guide.",
        "MAVSDK Python Documentation.",
        "ChromaDB Documentation.",
        "阿里云 DashScope 通义千问与文本向量服务文档。",
        "和风天气开发服务文档。",
        "牧野项目 README、ARCHITECTURE、QUALITY_SCORE、PROJECT_MEMORY 与源代码。",
    ]
    for idx, ref in enumerate(references, start=1):
        p = paragraph(document, f"[{idx}] {ref}", first_line=False)
        p.paragraph_format.first_line_indent = Cm(0)


def main() -> None:
    if not TEMPLATE.exists():
        raise FileNotFoundError(TEMPLATE)
    document = Document(TEMPLATE)
    clear_document(document)
    configure_document(document)
    add_report(document)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
