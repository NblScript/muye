"""Generate 牧野智农 presentation PPT."""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ── Colors ──
BG_CREAM = RGBColor(0xFA, 0xF8, 0xF5)
BG_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEXT_DARK = RGBColor(0x1D, 0x1B, 0x18)
TEXT_SECONDARY = RGBColor(0x6E, 0x68, 0x60)
TEXT_MUTED = RGBColor(0xA0, 0x98, 0x90)
ACCENT_AMBER = RGBColor(0xC2, 0x8B, 0x3A)
ACCENT_GREEN = RGBColor(0x3A, 0x9C, 0x6A)
ACCENT_PURPLE = RGBColor(0x7C, 0x6D, 0xCC)
ACCENT_BLUE = RGBColor(0x4A, 0x8F, 0xC2)
ACCENT_RED = RGBColor(0xC4, 0x50, 0x50)
ACCENT_CYAN = RGBColor(0x3A, 0x9C, 0xB0)
BORDER_LIGHT = RGBColor(0xE0, 0xDB, 0xD5)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
W = prs.slide_width
H = prs.slide_height


def set_slide_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_shape(slide, left, top, width, height, fill_color=None, border_color=None, border_width=Pt(1)):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color or BG_WHITE
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = border_width
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def add_text(slide, left, top, width, height, text, font_size=18, color=TEXT_DARK,
             bold=False, alignment=PP_ALIGN.LEFT, font_name='Microsoft YaHei'):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_bullet_list(slide, left, top, width, height, items, font_size=16, color=TEXT_DARK,
                    bullet_color=ACCENT_AMBER):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = 'Microsoft YaHei'
        p.space_after = Pt(8)
        p.level = 0
    return txBox


def add_tag(slide, left, top, text, color=ACCENT_AMBER):
    w, h = Inches(1.6), Inches(0.35)
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
    shape.fill.solid()
    r, g, b = int(str(color)[0:2], 16), int(str(color)[2:4], 16), int(str(color)[4:6], 16)
    shape.fill.fore_color.rgb = RGBColor(
        min(255, r + 180),
        min(255, g + 180),
        min(255, b + 180),
    )
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(10)
    p.font.color.rgb = color
    p.font.bold = True
    p.font.name = 'Microsoft YaHei'
    p.alignment = PP_ALIGN.CENTER
    tf.paragraphs[0].space_before = Pt(0)
    tf.paragraphs[0].space_after = Pt(0)
    return shape


def add_page_number(slide, num, total):
    add_text(slide, Inches(12.2), Inches(7.0), Inches(1), Inches(0.4),
             f'{num} / {total}', font_size=10, color=TEXT_MUTED, alignment=PP_ALIGN.RIGHT)


def add_section_header(slide, title, subtitle=''):
    add_shape(slide, Inches(0), Inches(0), W, H, BG_CREAM)
    # accent bar
    add_shape(slide, Inches(0), Inches(0), Inches(0.08), H, ACCENT_AMBER)
    # title
    add_text(slide, Inches(0.8), Inches(2.8), Inches(11), Inches(1),
             title, font_size=40, color=TEXT_DARK, bold=True)
    if subtitle:
        add_text(slide, Inches(0.8), Inches(3.9), Inches(11), Inches(0.8),
                 subtitle, font_size=20, color=TEXT_SECONDARY)


TOTAL_SLIDES = 11

# ════════════════════════════════════════
# Slide 1: Cover
# ════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
set_slide_bg(slide, BG_CREAM)

# accent stripe
add_shape(slide, Inches(0), Inches(0), Inches(0.12), H, ACCENT_AMBER)
add_shape(slide, Inches(0.12), Inches(0), Inches(0.04), H, ACCENT_GREEN)

# title block
add_text(slide, Inches(1.5), Inches(2.0), Inches(10), Inches(1.2),
         '牧野智农', font_size=56, color=TEXT_DARK, bold=True)
add_text(slide, Inches(1.5), Inches(3.2), Inches(10), Inches(0.8),
         '基于 YOLO + 千问大模型的智慧农业害虫防治闭环系统', font_size=22, color=TEXT_SECONDARY)

# tags
add_tag(slide, Inches(1.5), Inches(4.4), 'YOLOv8 检测', ACCENT_GREEN)
add_tag(slide, Inches(3.3), Inches(4.4), 'RAG 知识增强', ACCENT_PURPLE)
add_tag(slide, Inches(5.1), Inches(4.4), '千问大模型', ACCENT_BLUE)
add_tag(slide, Inches(6.9), Inches(4.4), 'PX4 无人机', ACCENT_AMBER)

add_text(slide, Inches(1.5), Inches(5.8), Inches(6), Inches(0.4),
         '害虫检测 → 气象采集 → AI 决策 → 无人机执行', font_size=14, color=TEXT_MUTED)
add_page_number(slide, 1, TOTAL_SLIDES)

# ════════════════════════════════════════
# Slide 2: Problem Background
# ════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, BG_CREAM)
add_shape(slide, Inches(0), Inches(0), Inches(0.08), H, ACCENT_AMBER)

add_text(slide, Inches(0.8), Inches(0.6), Inches(6), Inches(0.5),
         '问题背景', font_size=14, color=ACCENT_AMBER, bold=True)
add_text(slide, Inches(0.8), Inches(1.1), Inches(11), Inches(0.8),
         '农业害虫防治的三大痛点', font_size=36, color=TEXT_DARK, bold=True)

# three cards
cards = [
    ('识别滞后', '传统人工巡田效率低，\n发现虫害时往往已扩散，\n错过最佳防治窗口。', ACCENT_RED),
    ('决策盲目', '施药方案依赖经验，\n缺乏气象、农药知识和\n历史数据的综合分析。', ACCENT_AMBER),
    ('执行脱节', '识别、决策、喷洒各环节\n割裂，无法形成自动化\n闭环，人力成本高。', ACCENT_PURPLE),
]
for i, (title, desc, color) in enumerate(cards):
    x = Inches(0.8 + i * 4.0)
    y = Inches(2.5)
    card = add_shape(slide, x, y, Inches(3.6), Inches(3.8), BG_WHITE, BORDER_LIGHT)
    # color accent bar at top
    add_shape(slide, x, y, Inches(3.6), Inches(0.06), color)
    add_text(slide, x + Inches(0.3), y + Inches(0.4), Inches(3), Inches(0.5),
             title, font_size=22, color=color, bold=True)
    add_text(slide, x + Inches(0.3), y + Inches(1.2), Inches(3), Inches(2.2),
             desc, font_size=15, color=TEXT_SECONDARY)

add_page_number(slide, 2, TOTAL_SLIDES)

# ════════════════════════════════════════
# Slide 3: System Positioning
# ════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, BG_CREAM)
add_shape(slide, Inches(0), Inches(0), Inches(0.08), H, ACCENT_AMBER)

add_text(slide, Inches(0.8), Inches(0.6), Inches(6), Inches(0.5),
         '系统定位', font_size=14, color=ACCENT_AMBER, bold=True)
add_text(slide, Inches(0.8), Inches(1.1), Inches(11), Inches(0.8),
         '一句话说清楚', font_size=36, color=TEXT_DARK, bold=True)

# big quote
add_shape(slide, Inches(1.0), Inches(2.5), Inches(11.3), Inches(2.5), BG_WHITE, ACCENT_AMBER, Pt(2))
add_text(slide, Inches(1.5), Inches(2.8), Inches(10.3), Inches(1.8),
         '无人机航拍图片自动进入识别管线，发现虫情后联动\n气象数据与 RAG 农药知识库，由千问大模型生成施药方案，\n驱动无人机完成精准喷洒——全程无人值守。',
         font_size=22, color=TEXT_DARK)

# key metrics
metrics = [
    ('全自动', '文件监视 → 识别 → 决策 → 执行'),
    ('端到端', '从图像到喷洒的完整闭环'),
    ('可回放', '任务历史、决策过程全程记录'),
]
for i, (label, desc) in enumerate(metrics):
    x = Inches(1.0 + i * 3.8)
    y = Inches(5.5)
    add_shape(slide, x, y, Inches(3.4), Inches(1.2), BG_WHITE, BORDER_LIGHT)
    add_text(slide, x + Inches(0.3), y + Inches(0.15), Inches(2.8), Inches(0.4),
             label, font_size=18, color=ACCENT_AMBER, bold=True)
    add_text(slide, x + Inches(0.3), y + Inches(0.6), Inches(2.8), Inches(0.5),
             desc, font_size=13, color=TEXT_SECONDARY)

add_page_number(slide, 3, TOTAL_SLIDES)

# ════════════════════════════════════════
# Slide 4: Architecture
# ════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, BG_CREAM)
add_shape(slide, Inches(0), Inches(0), Inches(0.08), H, ACCENT_AMBER)

add_text(slide, Inches(0.8), Inches(0.6), Inches(6), Inches(0.5),
         '系统架构', font_size=14, color=ACCENT_AMBER, bold=True)
add_text(slide, Inches(0.8), Inches(1.1), Inches(11), Inches(0.8),
         '技术架构与数据流', font_size=36, color=TEXT_DARK, bold=True)

# pipeline flow - horizontal
stages = [
    ('无人机航拍', 'PX4 SITL\n定时采集农田图像', ACCENT_CYAN),
    ('图像监视', 'watchdog 文件系统监听\ndata/images/ 目录', ACCENT_BLUE),
    ('YOLO 检测', 'YOLOv8 本地推理\n害虫识别 + 置信度', ACCENT_GREEN),
    ('AI 决策', '气象 + RAG + 千问大模型\n生成施药方案', ACCENT_PURPLE),
    ('无人机执行', '航线规划 + PX4 飞控\n精准喷洒作业', ACCENT_AMBER),
]
for i, (title, desc, color) in enumerate(stages):
    x = Inches(0.5 + i * 2.5)
    y = Inches(2.4)
    card = add_shape(slide, x, y, Inches(2.2), Inches(2.4), BG_WHITE, BORDER_LIGHT)
    add_shape(slide, x, y, Inches(2.2), Inches(0.06), color)
    add_text(slide, x + Inches(0.2), y + Inches(0.3), Inches(1.8), Inches(0.4),
             title, font_size=15, color=color, bold=True, alignment=PP_ALIGN.CENTER)
    add_text(slide, x + Inches(0.15), y + Inches(0.9), Inches(1.9), Inches(1.2),
             desc, font_size=12, color=TEXT_SECONDARY, alignment=PP_ALIGN.CENTER)
    # arrow between cards
    if i < len(stages) - 1:
        add_text(slide, x + Inches(2.2), y + Inches(0.9), Inches(0.3), Inches(0.4),
                 '→', font_size=20, color=TEXT_MUTED, alignment=PP_ALIGN.CENTER)

# tech stack below
add_text(slide, Inches(0.8), Inches(5.2), Inches(4), Inches(0.4),
         '技术栈', font_size=14, color=ACCENT_AMBER, bold=True)

tech_items = [
    ('后端', 'Python + FastAPI + SQLAlchemy'),
    ('AI', 'YOLOv8 + 千问 (DashScope) + ChromaDB RAG'),
    ('前端', 'React + Vite + TypeScript'),
    ('仿真', 'PX4 SITL + MAVSDK'),
    ('存储', 'SQLite + JSONL 事件总线'),
]
for i, (label, value) in enumerate(tech_items):
    y = Inches(5.6) + Inches(i * 0.32)
    add_text(slide, Inches(0.8), y, Inches(1), Inches(0.3),
             label, font_size=11, color=ACCENT_AMBER, bold=True)
    add_text(slide, Inches(1.8), y, Inches(5), Inches(0.3),
             value, font_size=11, color=TEXT_SECONDARY)

add_page_number(slide, 4, TOTAL_SLIDES)

# ════════════════════════════════════════
# Slide 5: Detection Module
# ════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, BG_CREAM)
add_shape(slide, Inches(0), Inches(0), Inches(0.08), H, ACCENT_GREEN)

add_text(slide, Inches(0.8), Inches(0.6), Inches(6), Inches(0.5),
         '核心模块 ①', font_size=14, color=ACCENT_GREEN, bold=True)
add_text(slide, Inches(0.8), Inches(1.1), Inches(11), Inches(0.8),
         '害虫检测 — YOLOv8', font_size=36, color=TEXT_DARK, bold=True)

# left: description
add_shape(slide, Inches(0.8), Inches(2.4), Inches(5.5), Inches(4.5), BG_WHITE, BORDER_LIGHT)
items = [
    '基于 YOLOv8 的本地目标检测模型',
    '支持蚜虫、稻飞虱、粘虫等 10+ 种害虫',
    '图片到达后自动触发推理，无需人工干预',
    '返回害虫类型、置信度、位置（bounding box）',
    '检测结果自动持久化到 SQLite',
    '无虫害时标记"安全"，有虫害时触发后续管线',
]
for i, item in enumerate(items):
    y = Inches(2.8) + Inches(i * 0.55)
    add_text(slide, Inches(1.2), y, Inches(0.3), Inches(0.3),
             '●', font_size=8, color=ACCENT_GREEN)
    add_text(slide, Inches(1.6), y, Inches(4.5), Inches(0.5),
             item, font_size=14, color=TEXT_DARK)

# right: key metrics
add_shape(slide, Inches(7.0), Inches(2.4), Inches(5.5), Inches(4.5), BG_WHITE, BORDER_LIGHT)
add_shape(slide, Inches(7.0), Inches(2.4), Inches(5.5), Inches(0.06), ACCENT_GREEN)
add_text(slide, Inches(7.4), Inches(2.7), Inches(4.5), Inches(0.4),
         '关键指标', font_size=16, color=ACCENT_GREEN, bold=True)

metrics_data = [
    ('检测模型', 'YOLOv8 (ultralytics)'),
    ('权重文件', 'models/best.pt'),
    ('推理服务', '本地 API (端口 8010)'),
    ('支持害虫', '10+ 种（蚜虫、飞虱、螟虫等）'),
    ('输出格式', 'JSON (pest_type, confidence, bbox)'),
]
for i, (k, v) in enumerate(metrics_data):
    y = Inches(3.4) + Inches(i * 0.6)
    add_text(slide, Inches(7.4), y, Inches(2), Inches(0.3),
             k, font_size=12, color=TEXT_MUTED)
    add_text(slide, Inches(9.4), y, Inches(3), Inches(0.3),
             v, font_size=13, color=TEXT_DARK)

add_page_number(slide, 5, TOTAL_SLIDES)

# ════════════════════════════════════════
# Slide 6: Decision Module
# ════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, BG_CREAM)
add_shape(slide, Inches(0), Inches(0), Inches(0.08), H, ACCENT_PURPLE)

add_text(slide, Inches(0.8), Inches(0.6), Inches(6), Inches(0.5),
         '核心模块 ②', font_size=14, color=ACCENT_PURPLE, bold=True)
add_text(slide, Inches(0.8), Inches(1.1), Inches(11), Inches(0.8),
         'AI 决策 — 千问 + RAG + 气象', font_size=36, color=TEXT_DARK, bold=True)

# three pillars
pillars = [
    ('气象融合', '接入和风天气 API\n获取实时温度、湿度、\n风速等环境数据\n\n约束条件：风速 > 5 级\n或温度异常时暂停作业', ACCENT_CYAN),
    ('RAG 知识增强', 'ChromaDB 向量数据库\n存储农药目录、历史决策\n与农业知识文档\n\n检索相似案例注入提示词\n提升决策专业性', ACCENT_PURPLE),
    ('千问大模型', '调用 DashScope 千问 API\n结构化输出施药方案\n\n包含：农药名称、浓度、\n配比、飞行高度、\n喷洒速率、安全提示', ACCENT_AMBER),
]
for i, (title, desc, color) in enumerate(pillars):
    x = Inches(0.8 + i * 4.0)
    y = Inches(2.4)
    card = add_shape(slide, x, y, Inches(3.6), Inches(4.5), BG_WHITE, BORDER_LIGHT)
    add_shape(slide, x, y, Inches(3.6), Inches(0.06), color)
    add_text(slide, x + Inches(0.3), y + Inches(0.3), Inches(3), Inches(0.4),
             title, font_size=20, color=color, bold=True, alignment=PP_ALIGN.CENTER)
    add_text(slide, x + Inches(0.3), y + Inches(1.0), Inches(3), Inches(3.2),
             desc, font_size=14, color=TEXT_SECONDARY)

add_page_number(slide, 6, TOTAL_SLIDES)

# ════════════════════════════════════════
# Slide 7: Drone Module
# ════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, BG_CREAM)
add_shape(slide, Inches(0), Inches(0), Inches(0.08), H, ACCENT_AMBER)

add_text(slide, Inches(0.8), Inches(0.6), Inches(6), Inches(0.5),
         '核心模块 ③', font_size=14, color=ACCENT_AMBER, bold=True)
add_text(slide, Inches(0.8), Inches(1.1), Inches(11), Inches(0.8),
         '无人机执行 — PX4 航线规划与喷洒', font_size=36, color=TEXT_DARK, bold=True)

# left panel
add_shape(slide, Inches(0.8), Inches(2.4), Inches(5.5), Inches(4.5), BG_WHITE, BORDER_LIGHT)
items = [
    '根据地块边界自动生成覆盖航线',
    '固定走 PX4 SITL 工作流程',
    'PX4 SITL 模式下通过 MAVSDK 真实飞控',
    '执行前校验：高度、速度、喷洒速率、气象约束',
    '支持手动确认起飞 / 全自动起飞两种模式',
    '任务完成后自动记录喷洒面积、用时、轨迹',
]
for i, item in enumerate(items):
    y = Inches(2.8) + Inches(i * 0.55)
    add_text(slide, Inches(1.2), y, Inches(0.3), Inches(0.3),
             '●', font_size=8, color=ACCENT_AMBER)
    add_text(slide, Inches(1.6), y, Inches(4.5), Inches(0.5),
             item, font_size=14, color=TEXT_DARK)

# right: execution flow
add_shape(slide, Inches(7.0), Inches(2.4), Inches(5.5), Inches(4.5), BG_WHITE, BORDER_LIGHT)
add_shape(slide, Inches(7.0), Inches(2.4), Inches(5.5), Inches(0.06), ACCENT_AMBER)
add_text(slide, Inches(7.4), Inches(2.7), Inches(4.5), Inches(0.4),
         '执行流程', font_size=16, color=ACCENT_AMBER, bold=True)

flow_steps = [
    ('1. 航线规划', '根据地块 geofence 生成飞行路径'),
    ('2. 安全校验', '气象 / 高度 / 速度 / 电量约束检查'),
    ('3. 上传航线', '将航点上传至 PX4 飞控'),
    ('4. 起飞执行', '自动或手动确认后起飞'),
    ('5. 喷洒作业', '按航线飞行并执行喷洒'),
    ('6. 返航记录', '完成后返航并持久化喷洒记录'),
]
for i, (step, desc) in enumerate(flow_steps):
    y = Inches(3.4) + Inches(i * 0.55)
    add_text(slide, Inches(7.4), y, Inches(2.2), Inches(0.3),
             step, font_size=13, color=ACCENT_AMBER, bold=True)
    add_text(slide, Inches(9.6), y, Inches(2.8), Inches(0.3),
             desc, font_size=12, color=TEXT_SECONDARY)

add_page_number(slide, 7, TOTAL_SLIDES)

# ════════════════════════════════════════
# Slide 8: Frontend Dashboard
# ════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, BG_CREAM)
add_shape(slide, Inches(0), Inches(0), Inches(0.08), H, ACCENT_CYAN)

add_text(slide, Inches(0.8), Inches(0.6), Inches(6), Inches(0.5),
         '前端展示', font_size=14, color=ACCENT_CYAN, bold=True)
add_text(slide, Inches(0.8), Inches(1.1), Inches(11), Inches(0.8),
         '实时指挥大屏 — 三条视角同时展示', font_size=36, color=TEXT_DARK, bold=True)

# three view cards
views = [
    ('地图视角', 'SVG 实时地图\n地块边界、航线规划\n无人机位置、检测点位\n动画模拟飞行轨迹', ACCENT_GREEN),
    ('工作流视角', 'Pipeline 阶步进器\n检测→气象→决策→执行\n实时状态推送\n事件时间线', ACCENT_PURPLE),
    ('任务视角', '任务历史列表\n识别摘要、施药方案\n气象数据、安全提示\n支持筛选与搜索', ACCENT_AMBER),
]
for i, (title, desc, color) in enumerate(views):
    x = Inches(0.8 + i * 4.0)
    y = Inches(2.4)
    card = add_shape(slide, x, y, Inches(3.6), Inches(3.0), BG_WHITE, BORDER_LIGHT)
    add_shape(slide, x, y, Inches(3.6), Inches(0.06), color)
    add_text(slide, x + Inches(0.3), y + Inches(0.3), Inches(3), Inches(0.4),
             title, font_size=20, color=color, bold=True, alignment=PP_ALIGN.CENTER)
    add_text(slide, x + Inches(0.3), y + Inches(1.0), Inches(3), Inches(1.8),
             desc, font_size=14, color=TEXT_SECONDARY)

# tech note
add_shape(slide, Inches(0.8), Inches(5.8), Inches(11.7), Inches(1.0), BG_WHITE, BORDER_LIGHT)
add_text(slide, Inches(1.2), Inches(5.95), Inches(11), Inches(0.6),
         'React + Vite + TypeScript · CSS 变量驱动主题 · WebSocket 实时推送 · 自定义 UI 组件库（无第三方 UI 框架依赖）',
         font_size=13, color=TEXT_MUTED, alignment=PP_ALIGN.CENTER)

add_page_number(slide, 8, TOTAL_SLIDES)

# ════════════════════════════════════════
# Slide 9: Live Demo Placeholder
# ════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, BG_CREAM)
add_shape(slide, Inches(0), Inches(0), Inches(0.08), H, ACCENT_AMBER)

add_text(slide, Inches(0.8), Inches(0.6), Inches(6), Inches(0.5),
         '现场演示', font_size=14, color=ACCENT_AMBER, bold=True)
add_text(slide, Inches(0.8), Inches(1.1), Inches(11), Inches(0.8),
         '实时演示 — 全自动闭环流程', font_size=36, color=TEXT_DARK, bold=True)

add_shape(slide, Inches(1.5), Inches(2.5), Inches(10.3), Inches(3.5), BG_WHITE, ACCENT_AMBER, Pt(2))
add_text(slide, Inches(2.0), Inches(3.0), Inches(9.3), Inches(0.6),
         '演示步骤', font_size=22, color=ACCENT_AMBER, bold=True, alignment=PP_ALIGN.CENTER)

demo_steps = [
    '① 展示空闲态 Dashboard — 系统监视中，等待图像输入',
    '② 注入一张巡检图像 — 触发全自动管线',
    '③ 观察 Pipeline 推进 — YOLO 检测 → 气象采集 → AI 决策 → 无人机执行',
    '④ 查看识别结果 — 害虫类型、施药方案、航线规划',
]
for i, step in enumerate(demo_steps):
    y = Inches(3.8) + Inches(i * 0.5)
    add_text(slide, Inches(2.5), y, Inches(8), Inches(0.4),
             step, font_size=16, color=TEXT_DARK)

add_page_number(slide, 9, TOTAL_SLIDES)

# ════════════════════════════════════════
# Slide 10: Highlights
# ════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, BG_CREAM)
add_shape(slide, Inches(0), Inches(0), Inches(0.08), H, ACCENT_GREEN)

add_text(slide, Inches(0.8), Inches(0.6), Inches(6), Inches(0.5),
         '技术亮点', font_size=14, color=ACCENT_GREEN, bold=True)
add_text(slide, Inches(0.8), Inches(1.1), Inches(11), Inches(0.8),
         '与同类作品的差异化', font_size=36, color=TEXT_DARK, bold=True)

highlights = [
    ('全自动闭环', '文件监视 → 识别 → 决策 → 执行，无需人工干预，\n真正意义上的"无人值守智慧植保"', ACCENT_GREEN),
    ('RAG 知识增强', '不是简单的"调大模型 API"，而是将农药目录、\n历史决策、农业文档向量化后检索注入提示词', ACCENT_PURPLE),
    ('PX4 执行闭环', '无人机执行固定走 PX4 SITL，\n连接、上传航线、起飞、喷洒和返航状态可追踪', ACCENT_AMBER),
    ('端到端可视化', '地图、工作流、任务历史三条视角同时展示，\n决策过程可回放、可追溯', ACCENT_CYAN),
]
for i, (title, desc, color) in enumerate(highlights):
    x = Inches(0.8) if i % 2 == 0 else Inches(6.9)
    y = Inches(2.4) + (i // 2) * Inches(2.3)
    card = add_shape(slide, x, y, Inches(5.7), Inches(1.9), BG_WHITE, BORDER_LIGHT)
    add_shape(slide, x, y, Inches(0.08), Inches(1.9), color)
    add_text(slide, x + Inches(0.4), y + Inches(0.2), Inches(5), Inches(0.4),
             title, font_size=18, color=color, bold=True)
    add_text(slide, x + Inches(0.4), y + Inches(0.7), Inches(5), Inches(1.0),
             desc, font_size=13, color=TEXT_SECONDARY)

add_page_number(slide, 10, TOTAL_SLIDES)

# ════════════════════════════════════════
# Slide 11: Future & Thanks
# ════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, BG_CREAM)
add_shape(slide, Inches(0), Inches(0), Inches(0.08), H, ACCENT_AMBER)

add_text(slide, Inches(0.8), Inches(0.6), Inches(6), Inches(0.5),
         '未来展望', font_size=14, color=ACCENT_AMBER, bold=True)
add_text(slide, Inches(0.8), Inches(1.1), Inches(11), Inches(0.8),
         '从演示到生产', font_size=36, color=TEXT_DARK, bold=True)

future_items = [
    ('接入真实无人机', '在 PX4 链路稳定后对接商用植保无人机 API'),
    ('多机协同', '支持多架无人机并行作业，自动分配地块和避障'),
    ('边缘部署', '模型量化后部署到 Jetson 等边缘设备，田间实时推理'),
    ('移动端适配', '开发平板/手机端操控界面，支持田间现场操作'),
]
for i, (title, desc) in enumerate(future_items):
    y = Inches(2.4) + Inches(i * 1.1)
    add_shape(slide, Inches(0.8), y, Inches(11.7), Inches(0.9), BG_WHITE, BORDER_LIGHT)
    add_text(slide, Inches(1.2), y + Inches(0.1), Inches(3), Inches(0.35),
             title, font_size=16, color=ACCENT_AMBER, bold=True)
    add_text(slide, Inches(4.5), y + Inches(0.1), Inches(7.5), Inches(0.7),
             desc, font_size=14, color=TEXT_SECONDARY)

# thanks
add_text(slide, Inches(0.8), Inches(6.2), Inches(11.7), Inches(0.8),
         '感谢评审 · 欢迎提问', font_size=28, color=TEXT_DARK, bold=True, alignment=PP_ALIGN.CENTER)

add_page_number(slide, 11, TOTAL_SLIDES)

# ── Save ──
output_path = Path(__file__).resolve().parent.parent / 'docs' / '牧野智农-答辩PPT.pptx'
prs.save(output_path)
print(f'Saved to {output_path}')
